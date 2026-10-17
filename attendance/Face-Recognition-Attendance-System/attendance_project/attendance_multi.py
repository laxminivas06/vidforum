import os
import sys
import json
import time
import argparse
from datetime import datetime
import numpy as np
import cv2
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

from ai_face_engine import AIFaceEngine

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

STUDENTS_FOLDER = "students"
DATA_FILE = "students_data.json"
EXCEL_FILE = "attendance.xlsx"
JSON_OUTPUT = "attendance_latest.json"

def load_students():
    if not os.path.exists(DATA_FILE):
        return {}
    with open(DATA_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)
    normalized = {}
    for sid, info in data.items():
        if isinstance(info, str):
            normalized[sid] = {"name": info, "email": "", "class": "10", "section": "A"}
        else:
            normalized[sid] = info
    return normalized

def setup_excel(students):
    today = datetime.now().strftime("%Y-%m-%d")
    if os.path.exists(EXCEL_FILE):
        try:
            wb = openpyxl.load_workbook(EXCEL_FILE)
        except Exception:
            wb = openpyxl.Workbook()
            if "Sheet" in wb.sheetnames:
                del wb["Sheet"]
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    if today not in wb.sheetnames:
        ws = wb.create_sheet(title=today)
        headers = ["Student ID", "Name", "Class & Section", "Email", "Date", "Time", "Status", "Confidence"]
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        ws.column_dimensions['A'].width = 16
        ws.column_dimensions['B'].width = 24
        ws.column_dimensions['C'].width = 18
        ws.column_dimensions['D'].width = 28
        ws.column_dimensions['E'].width = 14
        ws.column_dimensions['F'].width = 14
        ws.column_dimensions['G'].width = 14
        ws.column_dimensions['H'].width = 16

        absent_fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
        row = 2
        for sid, info in students.items():
            name = info.get("name", sid)
            cls_sec = f"{info.get('class', '')} {info.get('section', '')}".strip()
            email = info.get("email", "")
            ws.cell(row=row, column=1, value=sid)
            ws.cell(row=row, column=2, value=name)
            ws.cell(row=row, column=3, value=cls_sec)
            ws.cell(row=row, column=4, value=email)
            ws.cell(row=row, column=5, value=today)
            ws.cell(row=row, column=6, value="--")
            
            sc = ws.cell(row=row, column=7, value="Absent")
            sc.fill = absent_fill
            sc.font = Font(color="CC0000", bold=True)
            sc.alignment = Alignment(horizontal="center", vertical="center")
            
            ws.cell(row=row, column=8, value="--").alignment = Alignment(horizontal="center", vertical="center")
            row += 1
        wb.save(EXCEL_FILE)
    return wb, today

def mark_present(wb, today, student_id, confidence=1.0):
    ws = wb[today]
    now_time = datetime.now().strftime("%I:%M:%S %p")
    present_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == str(student_id):
            if row[6].value != "Present":
                row[5].value = now_time
                row[6].value = "Present"
                row[6].fill = present_fill
                row[6].font = Font(color="006100", bold=True)
                row[6].alignment = Alignment(horizontal="center", vertical="center")
                row[7].value = f"{confidence * 100:.1f}%"
                row[7].alignment = Alignment(horizontal="center", vertical="center")
                wb.save(EXCEL_FILE)
                return True
            return False
    return False

def create_multi_student_frame(student_ids, target_w=960, target_h=540):
    """Combine multiple student portraits into a single multi-face classroom camera frame."""
    canvas = np.zeros((target_h, target_w, 3), dtype=np.uint8)
    canvas[:] = (25, 28, 38)
    
    n = len(student_ids)
    if n == 0:
        return canvas
        
    slot_w = target_w // n
    for i, sid in enumerate(student_ids):
        path = os.path.join(STUDENTS_FOLDER, f"{sid}.jpg")
        if not os.path.exists(path):
            continue
        img = cv2.imread(path)
        if img is None:
            continue
        h, w = img.shape[:2]
        scale = min((target_h - 100) / h, (slot_w - 30) / w)
        nw, nh = int(w * scale), int(h * scale)
        scaled = cv2.resize(img, (nw, nh), interpolation=cv2.INTER_CUBIC)
        
        ox = i * slot_w + (slot_w - nw) // 2
        oy = 60 + (target_h - 100 - nh) // 2
        canvas[oy:oy+nh, ox:ox+nw] = scaled
    return canvas

def process_frame_multi(frame, engine, wb, today, marked_students, students):
    """Detect and recognize all faces in frame simultaneously."""
    faces = engine.detect_faces(frame)
    matches = []
    
    for (x, y, fw, fh, conf) in faces:
        crop = frame[y:y+fh, x:x+fw]
        sid, name, score = engine.match_face(crop)
        matches.append((x, y, fw, fh, sid, name, score))
        
        if sid and sid not in marked_students:
            mark_present(wb, today, sid, score)
            marked_students[sid] = {
                "name": name,
                "time": datetime.now().strftime("%I:%M:%S %p"),
                "confidence": round(score * 100, 1)
            }
            print(f"[MULTI-FACE DETECTED] Present: {name} ({sid}) | Conf: {score*100:.1f}%")

    return matches

def draw_multi_hud(frame, matches, marked_students, total_students, title="MULTI-FACE ATTENDANCE"):
    h, w = frame.shape[:2]
    
    # Header
    cv2.rectangle(frame, (0, 0), (w, 50), (18, 22, 30), cv2.FILLED)
    cv2.putText(frame, f"VID SCHOOL | {title}", (15, 32),
                cv2.FONT_HERSHEY_DUPLEX, 0.65, (255, 255, 255), 1)
    
    stat_str = f"In Frame: {len(matches)} | Total Present: {len(marked_students)}/{total_students}"
    cv2.putText(frame, stat_str, (w - 380, 32),
                cv2.FONT_HERSHEY_SIMPLEX, 0.52, (0, 255, 200), 1)
                
    # Draw all detected faces
    for (x, y, fw, fh, sid, name, score) in matches:
        if sid:
            box_col = (0, 255, 120)
            tag = f"{name} ({sid})"
            sub = f"{score*100:.1f}%"
        else:
            box_col = (0, 0, 255)
            tag = "Unknown"
            sub = f"{score*100:.1f}%"
            
        cv2.rectangle(frame, (x, y), (x+fw, y+fh), box_col, 2)
        cv2.rectangle(frame, (x, max(0, y-34)), (x+max(fw, 140), y), box_col, cv2.FILLED)
        cv2.putText(frame, tag, (x+4, y-18),
                    cv2.FONT_HERSHEY_DUPLEX, 0.42, (0, 0, 0), 1)
        cv2.putText(frame, sub, (x+4, y-4),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.38, (30, 30, 30), 1)

    # Footer
    cv2.rectangle(frame, (0, h-25), (w, h), (18, 22, 30), cv2.FILLED)
    cv2.putText(frame, "Multi-Face AI Engine: YuNet Detection + MobileNetV3 Embeddings | Press 'Q' to exit", (15, h-7),
                cv2.FONT_HERSHEY_SIMPLEX, 0.4, (160, 160, 160), 1)
    return frame

def run_multi_demo(engine, students, wb, today, headless=False):
    """Demonstrate multiple students recognized simultaneously in the same frame."""
    print("\n" + "="*60)
    print("      VID SCHOOL - MULTI-FACE CONCURRENT ATTENDANCE DEMO")
    print("      Detecting multiple student faces in a single frame")
    print("="*60)
    
    marked_students = {}
    
    # Group students into sets of 3 for multi-face scenes
    student_keys = list(students.keys())
    groups = [
        student_keys[0:3],   # 10-A-01, 10-A-02, 10-A-03
        student_keys[3:6],   # 10-A-04, 10-A-05, 10-B-01
        student_keys[6:9],   # 10-B-02, 10-B-03, 09-A-01
        student_keys[9:10]   # 09-A-02
    ]
    
    if not headless:
        cv2.namedWindow("VID School - Multi-Face Attendance Demo", cv2.WINDOW_NORMAL)
        cv2.resizeWindow("VID School - Multi-Face Attendance Demo", 960, 540)
        
    for g_idx, group in enumerate(groups, 1):
        print(f"\n[*] Scene {g_idx}/{len(groups)}: Processing {len(group)} students simultaneously in frame: {', '.join(group)}")
        multi_frame = create_multi_student_frame(group, 960, 540)
        
        matches = process_frame_multi(multi_frame, engine, wb, today, marked_students, students)
        
        display = draw_multi_hud(multi_frame.copy(), matches, marked_students, len(students),
                                 title=f"SCENE {g_idx}: {len(matches)} STUDENTS SIMULTANEOUSLY")
        
        if not headless:
            cv2.imshow("VID School - Multi-Face Attendance Demo", display)
            key = cv2.waitKey(1500) & 0xFF
            if key == ord('q'):
                break
        else:
            time.sleep(0.05)

    if not headless:
        cv2.destroyAllWindows()

    print("\n" + "="*60)
    print(f"[*] Multi-Face Demo Complete! Marked {len(marked_students)}/{len(students)} students.")
    print(f"[+] Output Excel file: {os.path.abspath(EXCEL_FILE)}")
    print("="*60 + "\n")

def main():
    parser = argparse.ArgumentParser(description="Multi-Face Attendance Tracking")
    parser.add_argument("--demo", action="store_true", help="Run multi-face scene demo")
    parser.add_argument("--camera", action="store_true", help="Run live webcam multi-face tracking")
    parser.add_argument("--headless", action="store_true", help="Run without GUI")
    args = parser.parse_args()

    students = load_students()
    engine = AIFaceEngine()
    wb, today = setup_excel(students)

    if args.camera:
        print("[*] Starting camera in multi-face mode...")
        cap = cv2.VideoCapture(0)
        if not cap.isOpened():
            print("[-] Webcam not available. Falling back to multi-face demo.")
            run_multi_demo(engine, students, wb, today, headless=args.headless)
            return

        cv2.namedWindow("VID School - Multi-Face Camera", cv2.WINDOW_NORMAL)
        marked = {}
        while True:
            ret, frame = cap.read()
            if not ret:
                break
            matches = process_frame_multi(frame, engine, wb, today, marked, students)
            display = draw_multi_hud(frame.copy(), matches, marked, len(students))
            cv2.imshow("VID School - Multi-Face Camera", display)
            if cv2.waitKey(1) & 0xFF == ord('q'):
                break
        cap.release()
        cv2.destroyAllWindows()
    else:
        run_multi_demo(engine, students, wb, today, headless=args.headless)

if __name__ == "__main__":
    main()