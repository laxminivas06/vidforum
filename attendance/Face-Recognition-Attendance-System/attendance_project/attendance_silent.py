import cv2
import os
import json
import numpy as np
import pyautogui
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from deepface import DeepFace

STUDENTS_FOLDER = "students"
DATA_FILE = "students_data.json"
EXCEL_FILE = "attendance.xlsx"

def load_students():
    if os.path.exists(DATA_FILE):
        with open(DATA_FILE, "r") as f:
            return json.load(f)
    return {}

def setup_excel(students):
    today = datetime.now().strftime("%Y-%m-%d")
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    if today not in wb.sheetnames:
        ws = wb.create_sheet(title=today)
        headers = ["Student ID", "Name", "Date", "Time", "Status"]
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True, size=12)
            cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions['A'].width = 15
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 15
        ws.column_dimensions['E'].width = 15

        absent_fill = PatternFill(start_color="FFE0E0", end_color="FFE0E0", fill_type="solid")
        row = 2
        for sid, name in students.items():
            ws.cell(row=row, column=1, value=sid)
            ws.cell(row=row, column=2, value=name)
            ws.cell(row=row, column=3, value=today)
            ws.cell(row=row, column=4, value="--")
            sc = ws.cell(row=row, column=5, value="Absent")
            sc.fill = absent_fill
            sc.font = Font(color="CC0000", bold=True)
            sc.alignment = Alignment(horizontal="center")
            row += 1
        wb.save(EXCEL_FILE)
        print(f"✅ Excel sheet created for {today}")
    return wb, today

def mark_present(wb, today, student_id):
    ws = wb[today]
    now = datetime.now().strftime("%I:%M %p")
    present_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    for row in ws.iter_rows(min_row=2):
        if row[0].value == student_id:
            if row[4].value != "Present":
                row[3].value = now
                row[4].value = "Present"
                row[4].fill = present_fill
                row[4].font = Font(color="006100", bold=True)
                row[4].alignment = Alignment(horizontal="center")
                wb.save(EXCEL_FILE)
                return True
    return False

def run_attendance():
    students = load_students()
    if not students:
        print("❌ No students registered! Run register.py first.")
        return

    student_photos = {}
    for sid, name in students.items():
        photo_path = os.path.join(STUDENTS_FOLDER, f"{sid}.jpg")
        if os.path.exists(photo_path):
            student_photos[sid] = {"name": name, "photo": photo_path}
            print(f"✅ Loaded: {name} ({sid})")
        else:
            print(f"⚠️ Photo missing for {sid}")

    wb, today = setup_excel(students)
    marked_today = set()

    print("\n===== ATTENDANCE SYSTEM RUNNING =====")
    print("Running silently - watching BlueStacks screen...")
    print("Press Ctrl+C to stop.\n")

    frame_count = 0

    while True:
        try:
            # Take screenshot
            screenshot = pyautogui.screenshot()
            frame = cv2.cvtColor(np.array(screenshot), cv2.COLOR_RGB2BGR)
            frame_count += 1

            # Process every 5 frames
            if frame_count % 5 == 0:
                temp_path = "C:/attendance_project/temp_frame.jpg"
                cv2.imwrite(temp_path, frame)
                print(f"🔍 Scanning frame {frame_count}...")

                for sid, data in student_photos.items():
                    if sid in marked_today:
                        continue
                    try:
                        result = DeepFace.verify(
                            img1_path=temp_path,
                            img2_path=data["photo"],
                            model_name="Facenet",
                            enforce_detection=False,
                            distance_metric="cosine",
                            threshold=0.5,
                            silent=True
                        )
                        print(f"   {data['name']}: distance={result['distance']:.3f} verified={result['verified']}")
                        if result["verified"]:
                            if mark_present(wb, today, sid):
                                marked_today.add(sid)
                                print(f"\n✅ PRESENT: {data['name']} ({sid}) at {datetime.now().strftime('%I:%M %p')}\n")
                    except Exception as e:
                        print(f"   Error: {e}")

                if os.path.exists(temp_path):
                    os.remove(temp_path)

                # All students marked
                if len(marked_today) == len(student_photos):
                    print("🎉 All students marked present!")
                    break

        except KeyboardInterrupt:
            break

    print(f"\n✅ Attendance saved to {EXCEL_FILE}")
    print(f"Total Present: {len(marked_today)}/{len(students)}")

if __name__ == "__main__":
    run_attendance()
